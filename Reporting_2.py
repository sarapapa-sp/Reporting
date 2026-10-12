
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter


class ExcelReportGenerator:

    COLORS = {
        "green": "3C8D5A",
        "dark_grey": "56616B",
        "light_grey": "D9DDE2",
        "grand_total_green": "3C8D5A",
        "grand_total_blue": "4472C4",
        "level_1_fill": "E5E7E9",
        "level_1_border": "3C8D5A",
        "header_display_fill": "EAF4EE",
        "white": "FFFFFF",
        "text": "12263A",
        "navy": "17365D",
        "header_border": "7F8C8D",
    }

    GRAND_TOTAL_ROW_HEIGHT = 26.25
    LEVEL_1_ROW_HEIGHT = 26.25
    CHILD_ROW_HEIGHT = 18.75
    SPACER_ROW_HEIGHT = 2.25

    SECTION_HEADING_WIDTH = 8.5
    SMALL_SPACER_WIDTH = 0.5
    LABEL_COLUMN_WIDTH = 30
    VALUE_COLUMN_WIDTH = 16

    PERCENTAGE_MAX = 200
    PERCENTAGE_MIN = -100

    def __init__(
        self,
        file_name="global_markets_report.xlsx",
        sheet_name="Global Markets Report"
    ):
        self.file_name = file_name
        self.sheet_name = sheet_name
        self.workbook = None
        self.worksheet = None
        self._create_file()

    # =========================================================
    # WORKBOOK
    # =========================================================

    def _create_file(self):
        self.workbook = Workbook()
        self.worksheet = self.workbook.active
        self.worksheet.title = self.sheet_name

        ws = self.worksheet
        ws.freeze_panes = None
        ws.sheet_view.showGridLines = False

        # Native Excel outline configuration.
        # Summary/parent rows appear above their detail rows.
        ws.sheet_properties.outlinePr.summaryBelow = False
        ws.sheet_properties.outlinePr.summaryRight = True
        ws.sheet_view.showOutlineSymbols = True

    # =========================================================
    # CREATE REPORT
    # =========================================================

    def create_report(
        self,
        report_data,
        column_config,
        header_text,
        report_title,
        section_label=None,
        show_grand_total=False,
        grand_total_color="green",
        show_column_headers=True,
        collapse_hierarchy=False,
        start_row=3,
        start_col=3
    ):
        self._validate_report_parameters(
            report_data=report_data,
            column_config=column_config,
            grand_total_color=grand_total_color
        )

        section_col = start_col
        label_col = start_col + 2
        value_start_col = start_col + 4

        column_layout = self._build_column_layout(
            column_config=column_config,
            value_start_col=value_start_col
        )

        # Write headings when requested.
        if show_column_headers:
            self._write_header(
                header_text=header_text,
                column_layout=column_layout,
                row=start_row,
                label_col=label_col
            )

            header_spacer_row = start_row + 1

            self._write_spacer_row(
                row=header_spacer_row,
                start_col=start_col,
                number_of_columns=len(column_layout) + 5
            )

            next_row = header_spacer_row + 1
        else:
            next_row = start_row

        first_section_row = None
        last_section_row = None

        # Grand Total is not a collapsible hierarchy node.
        if show_grand_total:
            grand_total_row = next_row

            next_row = self._write_node(
                node=report_data,
                level=0,
                row=grand_total_row,
                label_col=label_col,
                column_layout=column_layout,
                report_title=report_title,
                grand_total_color=grand_total_color,
                collapse_hierarchy=False,
                write_children=False
            )

            self._write_spacer_row(
                row=next_row,
                start_col=start_col,
                number_of_columns=len(column_layout) + 5
            )
            next_row += 1

        # Write the hierarchy.
        for child in report_data.get("child", []):
            if first_section_row is None:
                first_section_row = next_row

            next_row = self._write_node(
                node=child,
                level=1,
                row=next_row,
                label_col=label_col,
                column_layout=column_layout,
                report_title=report_title,
                grand_total_color=grand_total_color,
                collapse_hierarchy=collapse_hierarchy,
                write_children=True
            )

            last_section_row = next_row - 1

        # Vertical merged section label.
        if (
            section_label
            and first_section_row is not None
            and last_section_row is not None
        ):
            self._write_section_label(
                section_label=section_label,
                start_row=first_section_row,
                end_row=last_section_row,
                col=section_col,
                grand_total_color=grand_total_color
            )

        self._apply_layout(
            column_layout=column_layout,
            start_col=start_col,
            label_col=label_col
        )

        return next_row

    # =========================================================
    # VALIDATION
    # =========================================================

    def _validate_report_parameters(
        self,
        report_data,
        column_config,
        grand_total_color
    ):
        if grand_total_color not in ("green", "blue"):
            raise ValueError(
                "grand_total_color must be 'green' or 'blue'."
            )

        if not isinstance(report_data, dict):
            raise TypeError("report_data must be a dictionary.")

        if not isinstance(column_config, dict):
            raise TypeError("column_config must be a dictionary.")

        for key, config in column_config.items():
            if "mapping_name" not in config:
                raise ValueError(
                    f"Missing mapping_name for column '{key}'."
                )

            if "category" not in config:
                raise ValueError(
                    f"Missing category for column '{key}'."
                )

    # =========================================================
    # COLUMN LAYOUT
    # =========================================================

    def _build_column_layout(
        self,
        column_config,
        value_start_col
    ):
        layout = []
        previous_part = None
        current_col = value_start_col

        for key, config in column_config.items():
            current_part = config.get("part")

            if (
                previous_part is not None
                and current_part != previous_part
            ):
                layout.append({
                    "kind": "separator",
                    "column": current_col
                })
                current_col += 1

            layout.append({
                "kind": "value",
                "column": current_col,
                "key": key,
                "config": config
            })

            current_col += 1
            previous_part = current_part

        return layout

    # =========================================================
    # COLUMN HEADERS
    # =========================================================

    def _write_header(
        self,
        header_text,
        column_layout,
        row,
        label_col
    ):
        ws = self.worksheet

        cell = ws.cell(
            row=row,
            column=label_col,
            value=header_text
        )

        cell.font = Font(
            name="Aptos",
            size=11,
            bold=True,
            color=self.COLORS["navy"]
        )
        cell.alignment = Alignment(
            horizontal="left",
            vertical="center"
        )
        cell.border = Border()

        for item in column_layout:
            col = item["column"]

            if item["kind"] == "separator":
                ws.cell(row=row, column=col).border = Border()
                continue

            config = item["config"]
            cell = ws.cell(
                row=row,
                column=col,
                value=config["mapping_name"]
            )

            self._apply_column_heading_style(
                cell=cell,
                color_grading=config.get("color_grading"),
                column_type=config.get("type")
            )

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True
            )
            cell.border = self._get_thin_border(
                self.COLORS["header_border"]
            )

        ws.row_dimensions[row].height = 28

    def _apply_column_heading_style(
        self,
        cell,
        color_grading,
        column_type=None
    ):
        if color_grading:
            grading = color_grading.lower().replace(" ", "_")

            grading_styles = {
                "green": (
                    self.COLORS["green"],
                    self.COLORS["white"]
                ),
                "dark_grey": (
                    self.COLORS["dark_grey"],
                    self.COLORS["white"]
                ),
                "light_grey": (
                    self.COLORS["light_grey"],
                    self.COLORS["text"]
                ),
            }

            if grading not in grading_styles:
                raise ValueError(
                    f"Unsupported color_grading: {color_grading}"
                )

            fill, font_color = grading_styles[grading]

        else:
            type_styles = {
                "Actuals": (
                    self.COLORS["green"],
                    self.COLORS["white"]
                ),
                "Previous": (
                    self.COLORS["light_grey"],
                    self.COLORS["text"]
                ),
                "Budget": (
                    self.COLORS["dark_grey"],
                    self.COLORS["white"]
                ),
            }

            fill, font_color = type_styles.get(
                column_type,
                (self.COLORS["white"], self.COLORS["text"])
            )

        cell.fill = PatternFill(
            fill_type="solid",
            fgColor=fill
        )
        cell.font = Font(
            name="Aptos",
            size=11,
            bold=True,
            color=font_color
        )

    # =========================================================
    # SECTION LABEL
    # =========================================================

    def _write_section_label(
        self,
        section_label,
        start_row,
        end_row,
        col,
        grand_total_color
    ):
        ws = self.worksheet

        border_color = (
            self.COLORS["grand_total_green"]
            if grand_total_color == "green"
            else self.COLORS["grand_total_blue"]
        )

        side = Side(style="thin", color=border_color)

        for row in range(start_row, end_row + 1):
            cell = ws.cell(row=row, column=col)
            cell.fill = PatternFill(
                fill_type="solid",
                fgColor=self.COLORS["header_display_fill"]
            )
            cell.border = Border(left=side, right=side)

        ws.merge_cells(
            start_row=start_row,
            start_column=col,
            end_row=end_row,
            end_column=col
        )

        cell = ws.cell(row=start_row, column=col)
        cell.value = section_label
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor=self.COLORS["header_display_fill"]
        )
        cell.font = Font(
            name="Aptos",
            size=15,
            bold=True,
            color=border_color
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            textRotation=90,
            wrap_text=True
        )
        cell.border = Border(left=side, right=side)

    # =========================================================
    # SPACER ROW
    # =========================================================

    def _write_spacer_row(
        self,
        row,
        start_col,
        number_of_columns
    ):
        ws = self.worksheet
        ws.row_dimensions[row].height = self.SPACER_ROW_HEIGHT

        for offset in range(number_of_columns):
            cell = ws.cell(
                row=row,
                column=start_col + offset
            )
            cell.value = None
            cell.fill = PatternFill(fill_type=None)
            cell.border = Border()

    # =========================================================
    # HIERARCHY NODE
    # =========================================================

    
    def _write_node(
        self,
        node,
        level,
        row,
        label_col,
        column_layout,
        report_title,
        grand_total_color,
        collapse_hierarchy,
        write_children=True
    ):
        ws = self.worksheet

        values = node.get("values", {})
        children = node.get("child", [])
        root_id = node.get("root_id", level)

        label = (
            report_title
            if root_id == 0
            else node.get("root_value", "")
        )

        # ---------------------------------------------------------
        # Label styling
        # ---------------------------------------------------------

        label_cell = ws.cell(
            row=row,
            column=label_col,
            value=label
        )

        if root_id == 0:
            font_size = 12
            bold = True
            italic = False
            font_color = self.COLORS["white"]

            fill = (
                self.COLORS["grand_total_green"]
                if grand_total_color == "green"
                else self.COLORS["grand_total_blue"]
            )

            label_cell.fill = PatternFill(
                fill_type="solid",
                fgColor=fill
            )
            label_cell.border = Border()
            ws.row_dimensions[row].height = self.GRAND_TOTAL_ROW_HEIGHT

        elif root_id == 1:
            font_size = 12
            bold = True
            italic = False
            font_color = self.COLORS["text"]

            label_cell.fill = PatternFill(
                fill_type="solid",
                fgColor=self.COLORS["level_1_fill"]
            )
            label_cell.border = self._get_bottom_border(
                self.COLORS["level_1_border"]
            )
            ws.row_dimensions[row].height = self.LEVEL_1_ROW_HEIGHT

        elif root_id == 2:
            font_size = 11
            bold = collapse_hierarchy
            italic = False
            font_color = self.COLORS["text"]

            label_cell.fill = PatternFill(
                fill_type="solid",
                fgColor=self.COLORS["white"]
            )
            label_cell.border = Border()
            ws.row_dimensions[row].height = self.CHILD_ROW_HEIGHT

        elif root_id == 3:
            font_size = 11
            bold = False
            italic = False
            font_color = self.COLORS["text"]

            label_cell.fill = PatternFill(
                fill_type="solid",
                fgColor=self.COLORS["white"]
            )
            label_cell.border = Border()
            ws.row_dimensions[row].height = self.CHILD_ROW_HEIGHT

        else:  # Root 4 and later
            font_size = 10
            bold = False
            italic = True
            font_color = self.COLORS["text"]

            label_cell.fill = PatternFill(
                fill_type="solid",
                fgColor=self.COLORS["white"]
            )
            label_cell.border = Border()
            ws.row_dimensions[row].height = self.CHILD_ROW_HEIGHT

        label_cell.font = Font(
            name="Aptos",
            size=font_size,
            bold=bold,
            italic=italic,
            color=font_color
        )

        label_cell.alignment = Alignment(
            horizontal="left" if root_id <= 1 else "right",
            vertical="center",
            indent=max(root_id, 0)
        )

        # ---------------------------------------------------------
        # Label/value spacer
        # ---------------------------------------------------------

        spacer = ws.cell(
            row=row,
            column=label_col + 1
        )
        spacer.value = None
        spacer.fill = PatternFill(
            fill_type="solid",
            fgColor=self.COLORS["white"]
        )
        spacer.border = Border()

        # ---------------------------------------------------------
        # Value columns
        # ---------------------------------------------------------

        for item in column_layout:
            col = item["column"]

            if item["kind"] == "separator":
                cell = ws.cell(row=row, column=col)
                cell.value = None
                cell.fill = PatternFill(
                    fill_type="solid",
                    fgColor=self.COLORS["white"]
                )
                cell.border = Border()
                continue

            key = item["key"]
            config = item["config"]

            value, number_format = self._format_value(
                values.get(key),
                config["category"]
            )

            cell = ws.cell(
                row=row,
                column=col,
                value=value
            )
            cell.number_format = number_format

            cell.alignment = Alignment(
                horizontal="right",
                vertical="center"
            )

            if root_id == 0:
                fill = (
                    self.COLORS["grand_total_green"]
                    if grand_total_color == "green"
                    else self.COLORS["grand_total_blue"]
                )
                cell.fill = PatternFill(
                    fill_type="solid",
                    fgColor=fill
                )
                cell.font = Font(
                    name="Aptos",
                    size=12,
                    bold=True,
                    color=self.COLORS["white"]
                )
                cell.border = Border()

            elif root_id == 1:
                cell.fill = PatternFill(
                    fill_type="solid",
                    fgColor=self.COLORS["level_1_fill"]
                )
                cell.font = Font(
                    name="Aptos",
                    size=12,
                    bold=True,
                    color=self.COLORS["text"]
                )
                cell.border = self._get_bottom_border(
                    self.COLORS["level_1_border"]
                )

            else:
                cell.fill = PatternFill(
                    fill_type="solid",
                    fgColor=self.COLORS["white"]
                )
                cell.font = Font(
                    name="Aptos",
                    size=font_size,
                    bold=False,
                    italic=italic,
                    color=self.COLORS["text"]
                )
                cell.border = Border()

        # ---------------------------------------------------------
        # Excel outline levels
        # ---------------------------------------------------------

        # A row with root_id 3 is outline level 1.
        # A row with root_id 4 is outline level 2, and so on.
        if root_id >= 3:
            ws.row_dimensions[row].outlineLevel = min(
                root_id - 2, 7
            )

        if not write_children or not children:
            return row + 1

        # ---------------------------------------------------------
        # Write descendants recursively
        # ---------------------------------------------------------

        child_start_row = row + 1
        next_row = child_start_row

        for child in children:
            next_row = self._write_node(
                node=child,
                level=level + 1,
                row=next_row,
                label_col=label_col,
                column_layout=column_layout,
                report_title=report_title,
                grand_total_color=grand_total_color,
                collapse_hierarchy=collapse_hierarchy,
                write_children=True
            )

        child_end_row = next_row - 1

        # ---------------------------------------------------------
        # Native Excel collapse
        # ---------------------------------------------------------

        # Root 2 and deeper nodes with children can be collapsed.
        # Root 0 and Root 1 are never collapsible.
        if collapse_hierarchy and root_id >= 2:
            ws.row_dimensions[row].collapsed = True

            # Hide the descendants initially. Their individual
            # outline levels are retained for nested Excel groups.
            for child_row in range(
                child_start_row,
                child_end_row + 1
            ):
                ws.row_dimensions[child_row].hidden = True

        return next_row

    # =========================================================
    # VALUE FORMATTING
    # =========================================================

    def _format_value(self, value, category):
        if value is None:
            return None, "General"

        if category == "sum":
            return value, '#,##0;[Red](#,##0);-'

        if category == "percentage":
            return self._parse_percentage(value)

        if category == "percentage-100":
            return self._parse_percentage(100 - value)

        return value, "General"

    def _parse_percentage(
        self,
        value,
        max_percentage=None,
        min_percentage=None
    ):
        if max_percentage is None:
            max_percentage = self.PERCENTAGE_MAX

        if min_percentage is None:
            min_percentage = self.PERCENTAGE_MIN

        if value > max_percentage:
            return f">{max_percentage}", "General"

        if value < min_percentage:
            return "NA", "General"

        return value / 100, '0.00%;[Red](0.00%);-'

    # =========================================================
    # BORDERS
    # =========================================================

    def _get_thin_border(self, color):
        side = Side(style="thin", color=color)
        return Border(
            left=side,
            right=side,
            top=side,
            bottom=side
        )

    def _get_bottom_border(self, color):
        side = Side(style="thin", color=color)
        return Border(bottom=side)

    # =========================================================
    # WORKSHEET LAYOUT
    # =========================================================

    def _apply_layout(
        self,
        column_layout,
        start_col,
        label_col
    ):
        ws = self.worksheet

        ws.column_dimensions[
            get_column_letter(start_col)
        ].width = self.SECTION_HEADING_WIDTH

        ws.column_dimensions[
            get_column_letter(start_col + 1)
        ].width = self.SMALL_SPACER_WIDTH

        ws.column_dimensions[
            get_column_letter(label_col)
        ].width = self.LABEL_COLUMN_WIDTH

        ws.column_dimensions[
            get_column_letter(label_col + 1)
        ].width = self.SMALL_SPACER_WIDTH

        for item in column_layout:
            col = item["column"]

            ws.column_dimensions[
                get_column_letter(col)
            ].width = (
                self.SMALL_SPACER_WIDTH
                if item["kind"] == "separator"
                else self.VALUE_COLUMN_WIDTH
            )

        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.freeze_panes = None

    # =========================================================
    # SAVE
    # =========================================================

    def save(self):
        self.workbook.save(self.file_name)
        return Path(self.file_name)
