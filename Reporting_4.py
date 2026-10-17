
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock, InlineFont
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter


class ExcelReportGenerator:

    COLORS = {
        "green": "3C8D5A",
        "dark_grey": "56616B",
        "light_grey": "D9DDE2",
        "grand_total_green": "3C8D5A",
        "grand_total_blue": "4472C4",
        "grand_total_purple": "7030A0",
        "level_1_fill": "E5E7E9",
        "level_1_border": "3C8D5A",
        "header_display_fill": "EAF4EE",
        "white": "FFFFFF",
        "arrow_green": "00A651",
        "arrow_red": "E53935",
        "text": "12263A",
        "navy": "17365D",
        "header_border": "7F8C8D",
    }

    GRAND_TOTAL_ROW_HEIGHT = 26.25
    LEVEL_1_ROW_HEIGHT = 26.25
    CHILD_ROW_HEIGHT = 18.75
    SPACER_ROW_HEIGHT = 2.25

    SECTION_HEADING_WIDTH = 8.5
    SMALL_SPACER_WIDTH = 0.5
    LABEL_COLUMN_WIDTH = 30
    VALUE_COLUMN_WIDTH = 16

    # Default percentage thresholds
    PERCENTAGE_MAX = 200
    PERCENTAGE_MIN = -100

    # Number formats
    NORMAL_PERCENT_FORMAT = '0.00%;[Red](0.00%);-'
    SIGNED_PERCENT_FORMAT = '+0.00%;-0.00%;-'

    def __init__(
        self,
        file_name="global_markets_report.xlsx",
        sheet_name="Global"
    ):
        self.file_name = file_name
        self.sheet_name = sheet_name
        self.workbook = None
        self.worksheet = None

        self._create_file()

    # =========================================================
    # WORKBOOK AND WORKSHEETS
    # =========================================================

    def _create_file(self):
        self.workbook = Workbook()
        self.worksheet = self.workbook.active
        self.worksheet.title = self.sheet_name

        self._configure_worksheet(self.worksheet)

    def _configure_worksheet(self, ws):
        ws.freeze_panes = None
        ws.sheet_view.showGridLines = False

        ws.sheet_properties.outlinePr.summaryBelow = False
        ws.sheet_properties.outlinePr.summaryRight = True
        ws.sheet_view.showOutlineSymbols = True

    def _select_worksheet(self, sheet_name):
        if sheet_name in self.workbook.sheetnames:
            ws = self.workbook[sheet_name]
        else:
            ws = self.workbook.create_sheet(title=sheet_name)
            self._configure_worksheet(ws)

        self.worksheet = ws
        return ws

    # =========================================================
    # CREATE REPORT
    # =========================================================

    def create_report(
        self,
        report_data,
        column_config,
        header_text,
        report_title,
        section_label=None,
        show_grand_total=False,
        grand_total_color="green",
        show_column_headers=True,
        collapse_hierarchy=False,
        start_row=3,
        start_col=3,
        sheet_name=None,
        collapse_from_root=1,
        percentage_min=None,
        percentage_max=None
    ):
        """
        Create a report in the selected worksheet.

        percentage_min / percentage_max:
            Optional report-specific percentage thresholds.
            Defaults to PERCENTAGE_MIN / PERCENTAGE_MAX.

        Supported percentage categories:
            percentage
            percentage_show
            percentage_100
            percentage_100_show
        """

        self._validate_report_parameters(
            report_data=report_data,
            column_config=column_config,
            grand_total_color=grand_total_color,
            collapse_from_root=collapse_from_root
        )

        # Resolve report-specific thresholds.
        if percentage_min is None:
            percentage_min = self.PERCENTAGE_MIN

        if percentage_max is None:
            percentage_max = self.PERCENTAGE_MAX

        if percentage_min >= percentage_max:
            raise ValueError(
                "percentage_min must be less than percentage_max."
            )

        if sheet_name is not None:
            self._select_worksheet(sheet_name)

        ws = self.worksheet
        outline_groups = []

        section_col = start_col
        label_col = start_col + 2
        value_start_col = start_col + 4

        column_layout = self._build_column_layout(
            column_config=column_config,
            value_start_col=value_start_col
        )

        # -----------------------------------------------------
        # Headers
        # -----------------------------------------------------

        if show_column_headers:
            self._write_header(
                header_text=header_text,
                column_layout=column_layout,
                row=start_row,
                label_col=label_col
            )

            header_spacer_row = start_row + 1

            self._write_spacer_row(
                row=header_spacer_row,
                start_col=start_col,
                number_of_columns=len(column_layout) + 5
            )

            next_row = header_spacer_row + 1
        else:
            next_row = start_row

        first_section_row = None
        last_section_row = None

        # -----------------------------------------------------
        # Grand Total
        # -----------------------------------------------------

        if show_grand_total:
            grand_total_row = next_row

            next_row = self._write_node(
                node=report_data,
                level=0,
                row=grand_total_row,
                label_col=label_col,
                column_layout=column_layout,
                report_title=report_title,
                grand_total_color=grand_total_color,
                collapse_hierarchy=collapse_hierarchy,
                collapse_from_root=collapse_from_root,
                outline_groups=outline_groups,
                percentage_min=percentage_min,
                percentage_max=percentage_max,
                write_children=False
            )

            self._write_spacer_row(
                row=next_row,
                start_col=start_col,
                number_of_columns=len(column_layout) + 5
            )

            next_row += 1

        # -----------------------------------------------------
        # Hierarchy
        # -----------------------------------------------------

        for child in report_data.get("child", []):
            if first_section_row is None:
                first_section_row = next_row

            next_row = self._write_node(
                node=child,
                level=1,
                row=next_row,
                label_col=label_col,
                column_layout=column_layout,
                report_title=report_title,
                grand_total_color=grand_total_color,
                collapse_hierarchy=collapse_hierarchy,
                collapse_from_root=collapse_from_root,
                outline_groups=outline_groups,
                percentage_min=percentage_min,
                percentage_max=percentage_max,
                write_children=True
            )

            last_section_row = next_row - 1

        # -----------------------------------------------------
        # Section label
        # -----------------------------------------------------

        if (
            section_label
            and first_section_row is not None
            and last_section_row is not None
        ):
            self._write_section_label(
                section_label=section_label,
                start_row=first_section_row,
                end_row=last_section_row,
                col=section_col,
                grand_total_color=grand_total_color
            )

        # -----------------------------------------------------
        # Excel outline groups
        # -----------------------------------------------------

        self._apply_outline_groups(
            outline_groups=outline_groups,
            collapse_hierarchy=collapse_hierarchy,
            collapse_from_root=collapse_from_root
        )

        self._apply_layout(
            column_layout=column_layout,
            start_col=start_col,
            label_col=label_col
        )

        return next_row

    # =========================================================
    # VALIDATION
    # =========================================================

    def _validate_report_parameters(
        self,
        report_data,
        column_config,
        grand_total_color,
        collapse_from_root
    ):
        if grand_total_color not in ("green", "blue", "purple"):
            raise ValueError(
                "grand_total_color must be 'green', 'blue', or 'purple'."
            )

        if not isinstance(report_data, dict):
            raise TypeError("report_data must be a dictionary.")

        if not isinstance(column_config, dict):
            raise TypeError("column_config must be a dictionary.")

        if (
            not isinstance(collapse_from_root, int)
            or isinstance(collapse_from_root, bool)
            or collapse_from_root < 1
        ):
            raise ValueError(
                "collapse_from_root must be an integer >= 1."
            )

        for key, config in column_config.items():
            if "mapping_name" not in config:
                raise ValueError(
                    f"Missing mapping_name for column '{key}'."
                )

            if "category" not in config:
                raise ValueError(
                    f"Missing category for column '{key}'."
                )

            supported_categories = {
                "sum",
                "percentage",
                "percentage_show",
                "percentage_100",
                "percentage_100_show",
                "percentage_arrow",
            }

            category = config["category"]

            if (
                category not in supported_categories
                and category is not None
            ):
                # Other categories are permitted and use General.
                pass

    # =========================================================
    # COLUMN LAYOUT
    # =========================================================

    def _build_column_layout(
        self,
        column_config,
        value_start_col
    ):
        layout = []
        previous_part = None
        current_col = value_start_col

        for key, config in column_config.items():
            current_part = config.get("part")

            if (
                previous_part is not None
                and current_part != previous_part
            ):
                layout.append({
                    "kind": "separator",
                    "column": current_col
                })
                current_col += 1

            layout.append({
                "kind": "value",
                "column": current_col,
                "key": key,
                "config": config
            })

            current_col += 1
            previous_part = current_part

        return layout

    # =========================================================
    # COLUMN HEADERS
    # =========================================================

    def _write_header(
        self,
        header_text,
        column_layout,
        row,
        label_col
    ):
        ws = self.worksheet

        cell = ws.cell(
            row=row,
            column=label_col,
            value=header_text
        )

        cell.font = Font(
            name="Aptos",
            size=11,
            bold=True,
            color=self.COLORS["navy"]
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )
        cell.border = Border()

        spacer = ws.cell(row=row, column=label_col + 1)
        spacer.value = None
        spacer.border = Border()

        for item in column_layout:
            col = item["column"]

            if item["kind"] == "separator":
                ws.cell(row=row, column=col).border = Border()
                continue

            config = item["config"]

            cell = ws.cell(
                row=row,
                column=col,
                value=config["mapping_name"]
            )

            self._apply_column_heading_style(
                cell=cell,
                color_grading=config.get("color_grading"),
                column_type=config.get("type")
            )

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True
            )
            cell.border = self._get_thin_border(
                self.COLORS["header_border"]
            )

        ws.row_dimensions[row].height = 28

    def _apply_column_heading_style(
        self,
        cell,
        color_grading,
        column_type=None
    ):
        if color_grading:
            grading = color_grading.lower().replace(" ", "_")

            styles = {
                "green": (
                    self.COLORS["green"],
                    self.COLORS["white"]
                ),
                "dark_grey": (
                    self.COLORS["dark_grey"],
                    self.COLORS["white"]
                ),
                "light_grey": (
                    self.COLORS["light_grey"],
                    self.COLORS["text"]
                )
            }

            if grading not in styles:
                raise ValueError(
                    f"Unsupported color_grading: {color_grading}"
                )

            fill, font_color = styles[grading]

        else:
            styles = {
                "Actuals": (
                    self.COLORS["green"],
                    self.COLORS["white"]
                ),
                "Previous": (
                    self.COLORS["light_grey"],
                    self.COLORS["text"]
                ),
                "Budget": (
                    self.COLORS["dark_grey"],
                    self.COLORS["white"]
                )
            }

            fill, font_color = styles.get(
                column_type,
                (self.COLORS["white"], self.COLORS["text"])
            )

        cell.fill = PatternFill(
            fill_type="solid",
            fgColor=fill
        )
        cell.font = Font(
            name="Aptos",
            size=11,
            bold=True,
            color=font_color
        )

    # =========================================================
    # SECTION LABEL
    # =========================================================

    def _write_section_label(
        self,
        section_label,
        start_row,
        end_row,
        col,
        grand_total_color
    ):
        ws = self.worksheet

        border_color = {
            "green": self.COLORS["grand_total_green"],
            "blue": self.COLORS["grand_total_blue"],
            "purple": self.COLORS["grand_total_purple"],
        }[grand_total_color]

        side = Side(style="thin", color=border_color)

        for row in range(start_row, end_row + 1):
            cell = ws.cell(row=row, column=col)

            cell.fill = PatternFill(
                fill_type="solid",
                fgColor=self.COLORS["header_display_fill"]
            )
            cell.border = Border(left=side, right=side)

        ws.merge_cells(
            start_row=start_row,
            start_column=col,
            end_row=end_row,
            end_column=col
        )

        cell = ws.cell(row=start_row, column=col)
        cell.value = section_label
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor=self.COLORS["header_display_fill"]
        )
        cell.font = Font(
            name="Aptos",
            size=15,
            bold=True,
            color=border_color
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            textRotation=90,
            wrap_text=True
        )
        cell.border = Border(left=side, right=side)

    # =========================================================
    # SPACER ROW
    # =========================================================

    def _write_spacer_row(
        self,
        row,
        start_col,
        number_of_columns
    ):
        ws = self.worksheet
        ws.row_dimensions[row].height = self.SPACER_ROW_HEIGHT

        for offset in range(number_of_columns):
            cell = ws.cell(
                row=row,
                column=start_col + offset
            )
            cell.value = None
            cell.fill = PatternFill(fill_type=None)
            cell.border = Border()

    # =========================================================
    # WRITE NODE
    # =========================================================

    def _write_node(
        self,
        node,
        level,
        row,
        label_col,
        column_layout,
        report_title,
        grand_total_color,
        collapse_hierarchy,
        collapse_from_root,
        outline_groups,
        percentage_min,
        percentage_max,
        write_children=True
    ):
        ws = self.worksheet

        values = node.get("values", {})
        children = node.get("child", [])
        root_id = node.get("root_id", level)

        label = (
            report_title
            if root_id == 0
            else node.get("root_value", "")
        )

        # -----------------------------------------------------
        # Label styling
        # -----------------------------------------------------

        label_cell = ws.cell(
            row=row,
            column=label_col,
            value=label
        )

        if root_id == 0:
            font_size = 12
            bold = True
            italic = False
            font_color = self.COLORS["white"]

            fill = {
                "green": self.COLORS["grand_total_green"],
                "blue": self.COLORS["grand_total_blue"],
                "purple": self.COLORS["grand_total_purple"],
            }[grand_total_color]

            label_cell.fill = PatternFill(
                fill_type="solid",
                fgColor=fill
            )
            label_cell.border = Border()
            ws.row_dimensions[row].height = self.GRAND_TOTAL_ROW_HEIGHT

        elif root_id == 1:
            font_size = 12
            bold = True
            italic = False
            font_color = self.COLORS["text"]

            label_cell.fill = PatternFill(
                fill_type="solid",
                fgColor=self.COLORS["level_1_fill"]
            )
            label_cell.border = self._get_bottom_border(
                self.COLORS["level_1_border"]
            )
            ws.row_dimensions[row].height = self.LEVEL_1_ROW_HEIGHT

        elif root_id == 2:
            font_size = 11
            bold = collapse_hierarchy
            italic = False
            font_color = self.COLORS["text"]

            label_cell.fill = PatternFill(
                fill_type="solid",
                fgColor=self.COLORS["white"]
            )
            label_cell.border = Border()
            ws.row_dimensions[row].height = self.CHILD_ROW_HEIGHT

        elif root_id == 3:
            font_size = 11
            bold = False
            italic = True
            font_color = self.COLORS["text"]

            label_cell.fill = PatternFill(
                fill_type="solid",
                fgColor=self.COLORS["white"]
            )
            label_cell.border = Border()
            ws.row_dimensions[row].height = self.CHILD_ROW_HEIGHT

        else:
            font_size = 10
            bold = False
            italic = True
            font_color = self.COLORS["text"]

            label_cell.fill = PatternFill(
                fill_type="solid",
                fgColor=self.COLORS["white"]
            )
            label_cell.border = Border()
            ws.row_dimensions[row].height = self.CHILD_ROW_HEIGHT

        label_cell.font = Font(
            name="Aptos",
            size=font_size,
            bold=bold,
            italic=italic,
            color=font_color
        )

        label_cell.alignment = Alignment(
            horizontal="left" if root_id <= 1 else "right",
            vertical="center",
            indent=max(root_id, 0)
        )

        # -----------------------------------------------------
        # Spacer
        # -----------------------------------------------------

        spacer = ws.cell(row=row, column=label_col + 1)
        spacer.value = None
        spacer.fill = PatternFill(
            fill_type="solid",
            fgColor=self.COLORS["white"]
        )
        spacer.border = Border()

        # -----------------------------------------------------
        # Values
        # -----------------------------------------------------

        for item in column_layout:
            col = item["column"]

            if item["kind"] == "separator":
                separator = ws.cell(row=row, column=col)
                separator.value = None
                separator.fill = PatternFill(
                    fill_type="solid",
                    fgColor=self.COLORS["white"]
                )
                separator.border = Border()
                continue

            key = item["key"]
            config = item["config"]

            raw_value = values.get(key)
            category = config["category"]

            value, number_format = self._format_value(
                value=raw_value,
                category=category,
                percentage_min=percentage_min,
                percentage_max=percentage_max
            )

            cell = ws.cell(row=row, column=col)
            if category == "percentage_arrow":
                cell.value = self._format_percentage_arrow(
                    value=raw_value,
                    percentage_min=percentage_min,
                    percentage_max=percentage_max,
                    value_color=self.COLORS["white"] if root_id == 0 else self.COLORS["text"]
                )
                cell.number_format = "General"
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.value = value
                cell.number_format = number_format
                cell.alignment = Alignment(horizontal="right", vertical="center")

            if root_id == 0:
                fill = {
                    "green": self.COLORS["grand_total_green"],
                    "blue": self.COLORS["grand_total_blue"],
                    "purple": self.COLORS["grand_total_purple"],
                }[grand_total_color]

                cell.fill = PatternFill(
                    fill_type="solid",
                    fgColor=fill
                )
                cell.font = Font(
                    name="Aptos",
                    size=12,
                    bold=True,
                    color=self.COLORS["white"]
                )
                cell.border = Border()

            elif root_id == 1:
                cell.fill = PatternFill(
                    fill_type="solid",
                    fgColor=self.COLORS["level_1_fill"]
                )
                cell.font = Font(
                    name="Aptos",
                    size=12,
                    bold=True,
                    color=self.COLORS["text"]
                )
                cell.border = self._get_bottom_border(
                    self.COLORS["level_1_border"]
                )

            else:
                cell.fill = PatternFill(
                    fill_type="solid",
                    fgColor=self.COLORS["white"]
                )
                cell.font = Font(
                    name="Aptos",
                    size=font_size,
                    bold=bold,
                    italic=italic,
                    color=self.COLORS["text"]
                )
                cell.border = Border()

        # -----------------------------------------------------
        # Children and outline metadata
        # -----------------------------------------------------

        if not write_children or not children:
            return row + 1

        child_start_row = row + 1
        next_row = child_start_row

        for child in children:
            next_row = self._write_node(
                node=child,
                level=level + 1,
                row=next_row,
                label_col=label_col,
                column_layout=column_layout,
                report_title=report_title,
                grand_total_color=grand_total_color,
                collapse_hierarchy=collapse_hierarchy,
                collapse_from_root=collapse_from_root,
                outline_groups=outline_groups,
                percentage_min=percentage_min,
                percentage_max=percentage_max,
                write_children=True
            )

        child_end_row = next_row - 1

        if root_id >= collapse_from_root:
            outline_groups.append({
                "root_id": root_id,
                "parent_row": row,
                "start_row": child_start_row,
                "end_row": child_end_row
            })

        return next_row

    # =========================================================
    # APPLY EXCEL OUTLINE GROUPS
    # =========================================================

    def _apply_outline_groups(
        self,
        outline_groups,
        collapse_hierarchy,
        collapse_from_root
    ):
        ws = self.worksheet

        # Process outer groups before inner groups.
        outline_groups.sort(
            key=lambda group: group["root_id"]
        )

        for group in outline_groups:
            root_id = group["root_id"]

            outline_level = min(
                root_id - collapse_from_root + 1,
                7
            )

            ws.row_dimensions.group(
                start=group["start_row"],
                end=group["end_row"],
                outline_level=outline_level,
                hidden=collapse_hierarchy
            )

            ws.row_dimensions[
                group["parent_row"]
            ].collapsed = collapse_hierarchy

    def _format_percentage_arrow(
        self,
        value,
        percentage_min=None,
        percentage_max=None,
        value_color=None
    ):
        """Return a colored arrow plus percentage as rich text in one Excel cell.

        Inputs use percentage units: 2.5 represents 2.5%, not 0.025.
        """
        if value is None or value == "":
            return "-"

        if percentage_min is None:
            percentage_min = self.PERCENTAGE_MIN
        if percentage_max is None:
            percentage_max = self.PERCENTAGE_MAX
        if value_color is None:
            value_color = self.COLORS["text"]

        try:
            numeric_value = float(value)
        except (TypeError, ValueError):
            return str(value)

        if numeric_value > percentage_max:
            return f">{percentage_max}"
        if numeric_value < percentage_min:
            return "NA"
        if numeric_value == 0:
            return "-"

        if numeric_value > 0:
            arrow = "▲ "
            arrow_color = self.COLORS["arrow_green"]
        else:
            arrow = "▼ "
            arrow_color = self.COLORS["arrow_red"]

        return CellRichText([
            TextBlock(InlineFont(color=arrow_color, b=True), arrow),
            TextBlock(InlineFont(color=value_color), f"{numeric_value:.2f}%")
        ])

    # =========================================================
    # VALUE FORMATTING
    # =========================================================

    def _format_value(
        self,
        value,
        category,
        percentage_min=None,
        percentage_max=None
    ):
        if value is None:
            return None, "General"

        if category == "sum":
            return value, '#,##0;[Red](#,##0);-'

        if category == "percentage":
            return self._parse_percentage(
                value=value,
                show_sign=False,
                subtract_from_100=False,
                percentage_min=percentage_min,
                percentage_max=percentage_max
            )

        if category == "percentage_show":
            return self._parse_percentage(
                value=value,
                show_sign=True,
                subtract_from_100=False,
                percentage_min=percentage_min,
                percentage_max=percentage_max
            )

        if category == "percentage_100":
            return self._parse_percentage(
                value=value,
                show_sign=False,
                subtract_from_100=True,
                percentage_min=percentage_min,
                percentage_max=percentage_max
            )

        if category == "percentage_100_show":
            return self._parse_percentage(
                value=value,
                show_sign=True,
                subtract_from_100=True,
                percentage_min=percentage_min,
                percentage_max=percentage_max
            )

        return value, "General"

    # =========================================================
    # PERCENTAGE PARSING
    # =========================================================

    def _parse_percentage(
        self,
        value,
        show_sign=False,
        subtract_from_100=False,
        percentage_min=None,
        percentage_max=None
    ):
        """
        Input percentage values are expressed as numbers such as
        25 for 25%, rather than 0.25.

        Thresholds are applied after the optional 100 - value
        transformation.
        """

        if percentage_min is None:
            percentage_min = self.PERCENTAGE_MIN

        if percentage_max is None:
            percentage_max = self.PERCENTAGE_MAX

        if subtract_from_100:
            value = 100 - value

        if value > percentage_max:
            return f">{percentage_max}", "General"

        if value < percentage_min:
            return "NA", "General"

        number_format = (
            self.SIGNED_PERCENT_FORMAT
            if show_sign
            else self.NORMAL_PERCENT_FORMAT
        )

        return value / 100, number_format

    # =========================================================
    # BORDERS
    # =========================================================

    def _get_thin_border(self, color):
        side = Side(style="thin", color=color)

        return Border(
            left=side,
            right=side,
            top=side,
            bottom=side
        )

    def _get_bottom_border(self, color):
        side = Side(style="thin", color=color)
        return Border(bottom=side)

    # =========================================================
    # WORKSHEET LAYOUT
    # =========================================================

    def _apply_layout(
        self,
        column_layout,
        start_col,
        label_col
    ):
        ws = self.worksheet

        ws.column_dimensions[
            get_column_letter(start_col)
        ].width = self.SECTION_HEADING_WIDTH

        ws.column_dimensions[
            get_column_letter(start_col + 1)
        ].width = self.SMALL_SPACER_WIDTH

        ws.column_dimensions[
            get_column_letter(label_col)
        ].width = self.LABEL_COLUMN_WIDTH

        ws.column_dimensions[
            get_column_letter(label_col + 1)
        ].width = self.SMALL_SPACER_WIDTH

        for item in column_layout:
            col = item["column"]

            ws.column_dimensions[
                get_column_letter(col)
            ].width = (
                self.SMALL_SPACER_WIDTH
                if item["kind"] == "separator"
                else self.VALUE_COLUMN_WIDTH
            )

        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.freeze_panes = None

    # =========================================================
    # SAVE
    # =========================================================

    def save(self):
        self.workbook.save(self.file_name)
        return Path(self.file_name)
